#!/usr/bin/env python3
"""
FABRIC ORDERS (PENDING) EXPORT ENGINE
Directly queries the Garment ERP database (Supabase PostgreSQL) to export
all Purchase Orders / items filtered by Status='Pending'.

Matches the exact data structure and output of:
https://garment-erp-wiwk.onrender.com/fabric_orders?tab=report (Status: Pending -> Export Excel)
"""

import os
import sys
import json
import datetime
from pathlib import Path

try:
    import psycopg2
except ImportError:
    psycopg2 = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

from dotenv import load_dotenv

DEFAULT_FILENAME = "Fabric Orders Pending.xlsx"
DEFAULT_DB_URL = "postgresql://postgres.cusbryojwnldabchhfkx:9788%40Srinithi@aws-1-ap-northeast-1.pooler.supabase.com:6543/postgres"


def order_item_process(order_dict, item):
    """Fallback order-level process values if item-level fields are missing."""
    return {
        'process_type': item.get('process_type') or order_dict.get('process_type') or '',
        'process': item.get('process') or order_dict.get('process') or '',
        'fabric_incharge': item.get('fabric_incharge') or order_dict.get('fabric_incharge') or '',
    }


def format_cell_list(val):
    """Converts string, list, or None into a clean comma-separated string."""
    if val is None:
        return ''
    if isinstance(val, (list, tuple)):
        return ', '.join(str(v).strip() for v in val if str(v).strip())
    return str(val).strip()


def run_fabric_orders_download(
    db_url=None,
    download_dir=None,
    status_filter="Pending",
    status_callback=None,
    progress_callback=None,
    stop_event=None
):
    """
    Exports Pending Fabric Orders into an Excel spreadsheet.

    Args:
        db_url (str): PostgreSQL connection string. Defaults to GARMENT_ERP_DATABASE_URL from .env
        download_dir (str|Path): Destination base folder. Defaults to DOWNLOAD_DIR from .env
        status_filter (str): Filter value for item status, defaults to 'Pending' (case-insensitive).
        status_callback (callable): Optional callback for status messages.
        progress_callback (callable): Optional callback(current, total).
        stop_event (threading.Event): Optional cancellation event.

    Returns:
        tuple (bool, str): (Success flag, Result message)
    """
    def log(msg):
        print(msg)
        if status_callback:
            try:
                status_callback(msg)
            except Exception:
                pass

    def set_progress(curr, total):
        if progress_callback:
            try:
                progress_callback(curr, total)
            except Exception:
                pass

    # 1. Load environment if needed
    env_path = Path(__file__).parent.resolve() / ".env"
    if env_path.exists():
        load_dotenv(dotenv_path=env_path)

    if not db_url:
        db_url = os.getenv("GARMENT_ERP_DATABASE_URL") or DEFAULT_DB_URL

    if not download_dir:
        download_dir = os.getenv("DOWNLOAD_DIR", r"C:\ERP_DOWNLOADS")

    dest_dir = Path(download_dir).resolve() / "Fabric Orders"
    dest_dir.mkdir(parents=True, exist_ok=True)

    if not psycopg2:
        err_msg = "psycopg2 library is not installed. Please run: pip install psycopg2-binary"
        log(f"[ERROR] {err_msg}")
        return False, err_msg

    if not Workbook:
        err_msg = "openpyxl library is not installed. Please run: pip install openpyxl"
        log(f"[ERROR] {err_msg}")
        return False, err_msg

    set_progress(0, 1)
    log(f"Connecting to Garment ERP Database...")
    log(f"  Target Destination: {dest_dir}")
    log(f"  Status Filter     : {status_filter}")

    conn = None
    try:
        conn = psycopg2.connect(db_url)
        cur = conn.cursor()

        log("  Executing query on 'generated_orders'...")
        cur.execute("""
            SELECT id, po_number, order_items, process_type, process, fabric_incharge, status, created_at
            FROM generated_orders
            ORDER BY created_at DESC;
        """)
        orders = cur.fetchall()
        log(f"  Retrieved {len(orders)} total Purchase Orders from database.")

        if stop_event and stop_event.is_set():
            log("Operation cancelled by user.")
            return False, "Cancelled"

        export_rows = []
        target_status_clean = str(status_filter).strip().lower()

        for order in orders:
            if stop_event and stop_event.is_set():
                log("Operation cancelled by user.")
                return False, "Cancelled"

            o_id, po_num, items_raw, def_proc_type, def_proc, def_incharge, o_status, created_at = order
            order_dict = {
                'po_number': po_num,
                'process_type': def_proc_type,
                'process': def_proc,
                'fabric_incharge': def_incharge,
                'status': o_status
            }

            try:
                items = json.loads(items_raw) if isinstance(items_raw, str) else (items_raw or [])
            except Exception:
                items = []

            for item in items:
                proc_info = order_item_process(order_dict, item)
                item_status = item.get('status') or o_status or 'Pending'
                
                # Filter by status (if target_status_clean is set)
                if target_status_clean:
                    if str(item_status).strip().lower() != target_status_clean:
                        continue

                created_str = created_at.isoformat() if created_at else ""

                export_rows.append([
                    po_num or '',
                    item.get('fabric_name', ''),
                    item.get('uom', ''),
                    item.get('gsm', ''),
                    format_cell_list(item.get('colour', [])),
                    format_cell_list(item.get('dia', [])),
                    item.get('order_qty', ''),
                    proc_info.get('process_type', ''),
                    proc_info.get('process', ''),
                    proc_info.get('fabric_incharge', ''),
                    item_status or '',
                    created_str
                ])

        log(f"  Found {len(export_rows)} matching item rows with Status = '{status_filter}'.")

        # Create Excel Workbook matching garment-erp format
        log("  Generating formatted Excel workbook...")
        wb = Workbook()
        ws = wb.active
        ws.title = "Fabric Orders"

        headers = [
            'PO No', 'Fabric', 'UOM', 'GSM', 'Colour', 'DIA',
            'Order Qty', 'Process Type', 'Process', 'Fabric Incharge', 'Status', 'Created At'
        ]
        ws.append(headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Add data rows
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        data_font = Font(name="Calibri", size=10)

        for row_idx, row in enumerate(export_rows, start=2):
            ws.append(row)
            for col_idx in range(1, len(row) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                # Right align quantity
                if col_idx == 7:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in (1, 3, 4, 11):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 38)

        # Freeze top header row
        ws.freeze_panes = "A2"

        # Save single Excel file
        target_path = dest_dir / DEFAULT_FILENAME
        wb.save(str(target_path))

        file_size_bytes = target_path.stat().st_size
        file_size_kb = file_size_bytes / 1024

        set_progress(1, 1)
        log("\n" + "=" * 60)
        log("FABRIC ORDERS (PENDING) EXPORT COMPLETED SUCCESSFULLY!")
        log(f"  Saved File  : {target_path}")
        log(f"  Total Rows  : {len(export_rows)}")
        log(f"  File Size   : {file_size_kb:.2f} KB ({file_size_bytes:,} bytes)")
        log("=" * 60)

        return True, f"Successfully exported {len(export_rows)} pending rows to {DEFAULT_FILENAME}"

    except Exception as e:
        err_str = f"Fabric Orders export error: {str(e)}"
        log(f"\n[FATAL ERROR] {err_str}")
        return False, err_str

    finally:
        if conn:
            try:
                conn.close()
            except Exception:
                pass


if __name__ == "__main__":
    from dotenv import load_dotenv
    load_dotenv()

    base_folder = os.getenv("DOWNLOAD_DIR", r"C:\ERP_DOWNLOADS")
    db_conn = os.getenv("GARMENT_ERP_DATABASE_URL")

    print("===================================================")
    print("  Garment ERP - Fabric Orders (Pending) Downloader")
    print("===================================================")
    print(f"Target Base Folder: {base_folder}")
    print(f"Destination: {Path(base_folder) / 'Fabric Orders'}")
    print("---------------------------------------------------")

    success, msg = run_fabric_orders_download(
        db_url=db_conn,
        download_dir=base_folder,
        status_filter="Pending"
    )

    if not success:
        sys.exit(1)
