import os
import sys
import json
import time
import datetime
import threading
from pathlib import Path
import pandas as pd
import openpyxl

AUTO_DOWNLOAD_DIR = Path(r"C:\Users\santhosh\Documents\old\auto download")
LOCAL_FILTER_CONFIG_FILE = Path(__file__).parent.resolve() / "config" / "mainout_filter_config.json"
SOURCE_FILTER_CONFIG_FILE = AUTO_DOWNLOAD_DIR / "mainout_filter_config.json"

LOCAL_FABRIC_LIST_FILE = Path(__file__).parent.resolve() / "config" / "fabric_list.json"
SOURCE_FABRIC_LIST_FILE = AUTO_DOWNLOAD_DIR / "fabric_list.json"
LOCAL_WIP_LIST_FILE = Path(__file__).parent.resolve() / "config" / "production_wip_list.json"
SOURCE_WIP_LIST_FILE = AUTO_DOWNLOAD_DIR / "production_wip_list.json"

DEFAULT_FABRIC_LIST = [
    {"fabric": "CRYSTAL LYCRA", "unit": "Kgs", "gsm": "", "enabled": True},
    {"fabric": "LYCRA JERSEY", "unit": "Kgs", "gsm": "185", "enabled": True},
    {"fabric": "LYCRA JERSEY", "unit": "Kgs", "gsm": "210", "enabled": True},
    {"fabric": "NYLON SHIMMER", "unit": "Kgs", "gsm": "", "enabled": True},
    {"fabric": "POLY SHIMMER", "unit": "Kgs", "gsm": "", "enabled": True},
    {"fabric": "VISCOSE LYCRA JERSEY", "unit": "Kgs", "gsm": "", "enabled": True},
    {"fabric": "CROSS LOOP DESIGN", "unit": "Kgs", "gsm": "", "enabled": True},
    {"fabric": "SINGLE JERSEY", "unit": "Kgs", "gsm": "165", "enabled": True},
    {"fabric": "METTALIC", "unit": "Mtr", "gsm": "", "enabled": True},
    {"fabric": "RAYON SLUB LYCRA BFOLD", "unit": "Mtr", "gsm": "", "enabled": True},
    {"fabric": "RAYON", "unit": "Mtr", "gsm": "", "enabled": True},
    {"fabric": "HAVY RAYON", "unit": "Mtr", "gsm": "", "enabled": True}
]

DEFAULT_WIP_LIST = [
    {"grouping": "CUTTING", "enabled": True},
    {"grouping": "STITCHING", "enabled": True},
    {"grouping": "SNG - HALF FINISHING", "enabled": True},
    {"grouping": "IRONING & PACKING", "enabled": True}
]

class ErpAutoSyncService:
    def __init__(self):
        self.lock = threading.Lock()
        self.stop_event = threading.Event()
        self.worker_thread = None
        self.job_state = {
            "status": "idle", # idle, running, completed, error, cancelled
            "mode": "review",  # review (Option A) or auto_save (Option B)
            "feed_strategy": "replace", # replace (delete old & save new) or append (add new)
            "step_index": 0,
            "total_steps": 8,
            "step_name": "",
            "progress_percent": 0,
            "logs": [],
            "error": None,
            "output_file": None,
            "summary": {},
            "started_at": None,
            "completed_at": None
        }

    def get_filter_config(self):
        """Loads sheet-wise whitelist filter config from local config or source."""
        if LOCAL_FILTER_CONFIG_FILE.exists():
            try:
                with open(LOCAL_FILTER_CONFIG_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception as e:
                print(f"[WARN] Error reading local filter config: {e}")

        if SOURCE_FILTER_CONFIG_FILE.exists():
            try:
                with open(SOURCE_FILTER_CONFIG_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.save_filter_config(data)
                    return data
            except Exception as e:
                print(f"[WARN] Error reading source filter config: {e}")

        return {
            "enabled": True,
            "filters": {
                "Fabric Stock": [],
                "Fabric WIP": [],
                "Production WIP": [],
                "Pending Orders": [],
                "Finished Goods": []
            }
        }

    def save_filter_config(self, config_data):
        """Saves sheet-wise whitelist filter configuration to local config file."""
        try:
            LOCAL_FILTER_CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
            with open(LOCAL_FILTER_CONFIG_FILE, "w", encoding="utf-8") as f:
                json.dump(config_data, f, indent=2, ensure_ascii=False)
            return True
        except Exception as e:
            print(f"[ERROR] Error saving filter config: {e}")
            return False

    def get_fabric_list(self):
        """Loads Fabric Stock items to download from local config or source."""
        if LOCAL_FABRIC_LIST_FILE.exists():
            try:
                with open(LOCAL_FABRIC_LIST_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception as e:
                print(f"[WARN] Error reading local fabric list: {e}")

        if SOURCE_FABRIC_LIST_FILE.exists():
            try:
                with open(SOURCE_FABRIC_LIST_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.save_fabric_list(data)
                    return data
            except Exception as e:
                print(f"[WARN] Error reading source fabric list: {e}")

        return list(DEFAULT_FABRIC_LIST)

    def save_fabric_list(self, items):
        """Saves Fabric Stock items to local config file."""
        try:
            LOCAL_FABRIC_LIST_FILE.parent.mkdir(parents=True, exist_ok=True)
            with open(LOCAL_FABRIC_LIST_FILE, "w", encoding="utf-8") as f:
                json.dump(items, f, indent=2, ensure_ascii=False)
            return True
        except Exception as e:
            print(f"[ERROR] Error saving fabric list: {e}")
            return False

    def reset_fabric_list(self):
        """Resets Fabric Stock items to default list."""
        self.save_fabric_list(DEFAULT_FABRIC_LIST)
        return list(DEFAULT_FABRIC_LIST)

    def get_wip_list(self):
        """Loads Production WIP groupings to download from local config or source."""
        if LOCAL_WIP_LIST_FILE.exists():
            try:
                with open(LOCAL_WIP_LIST_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception as e:
                print(f"[WARN] Error reading local WIP list: {e}")

        if SOURCE_WIP_LIST_FILE.exists():
            try:
                with open(SOURCE_WIP_LIST_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.save_wip_list(data)
                    return data
            except Exception as e:
                print(f"[WARN] Error reading source WIP list: {e}")

        return list(DEFAULT_WIP_LIST)

    def save_wip_list(self, items):
        """Saves Production WIP groupings to local config file."""
        try:
            LOCAL_WIP_LIST_FILE.parent.mkdir(parents=True, exist_ok=True)
            with open(LOCAL_WIP_LIST_FILE, "w", encoding="utf-8") as f:
                json.dump(items, f, indent=2, ensure_ascii=False)
            return True
        except Exception as e:
            print(f"[ERROR] Error saving WIP list: {e}")
            return False

    def reset_wip_list(self):
        """Resets Production WIP groupings to default list."""
        self.save_wip_list(DEFAULT_WIP_LIST)
        return list(DEFAULT_WIP_LIST)

    def _log(self, message):
        timestamp = datetime.datetime.now().strftime("%H:%M:%S")
        log_line = f"[{timestamp}] {message}"
        print(log_line)
        with self.lock:
            self.job_state["logs"].append(log_line)
            # Keep max 500 lines to prevent memory bloat
            if len(self.job_state["logs"]) > 500:
                self.job_state["logs"] = self.job_state["logs"][-500:]

    def get_status(self):
        with self.lock:
            return dict(self.job_state)

    def load_source_config(self):
        """Reads configuration and active item counts."""
        if not AUTO_DOWNLOAD_DIR.exists():
            return {
                "available": False,
                "error": f"Directory not found: {AUTO_DOWNLOAD_DIR}"
            }

        env_vars = {}
        env_file = AUTO_DOWNLOAD_DIR / ".env"
        if env_file.exists():
            try:
                with open(env_file, "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line and not line.startswith("#") and "=" in line:
                            k, v = line.split("=", 1)
                            env_vars[k.strip()] = v.strip().strip('"').strip("'")
            except Exception as e:
                print(f"[WARN] Error reading .env: {e}")

        username = env_vars.get("ERP_USERNAME", "")
        masked_user = (username[:3] + "***@" + username.split("@")[-1]) if "@" in username else "***"
        download_dir = env_vars.get("DOWNLOAD_DIR", r"C:\ERP_DOWNLOADS")

        # Load active item counts from local config
        fab_items = self.get_fabric_list()
        fabric_total = len(fab_items)
        fabric_active = len([x for x in fab_items if x.get("enabled", True)])

        wip_items = self.get_wip_list()
        wip_total = len(wip_items)
        wip_active = len([x for x in wip_items if x.get("enabled", True)])

        # Check latest mainout.xlsx
        output_candidate = Path(download_dir) / "output" / "mainout.xlsx"
        has_existing_mainout = output_candidate.exists()
        existing_info = None
        if has_existing_mainout:
            mtime = datetime.datetime.fromtimestamp(os.path.getmtime(output_candidate)).strftime("%Y-%m-%d %H:%M:%S")
            size_kb = os.path.getsize(output_candidate) / 1024
            existing_info = {
                "path": str(output_candidate),
                "modified": mtime,
                "size_kb": round(size_kb, 1)
            }

        return {
            "available": True,
            "source_dir": str(AUTO_DOWNLOAD_DIR),
            "username_masked": masked_user,
            "default_download_dir": download_dir,
            "fabric_count": fabric_active,
            "fabric_total": fabric_total,
            "wip_count": wip_active,
            "wip_total": wip_total,
            "has_existing_mainout": has_existing_mainout,
            "existing_mainout": existing_info
        }

    def stop(self):
        """Signals stop event to cancel currently running process."""
        self.stop_event.set()
        self._log("[CANCEL] User requested cancellation. Stopping after current operation...")

    def parse_mainout_file(self, file_path):
        """
        Parses mainout.xlsx into standard 5-sheet dictionary matching Bulk Feed UI expectations.
        """
        p = Path(file_path)
        if not p.exists():
            raise FileNotFoundError(f"File does not exist: {file_path}")

        expected_sheets = ['Fabric Stock', 'Fabric WIP', 'Production WIP', 'Pending Orders', 'Finished Goods']
        parsed_sheets = {s: [] for s in expected_sheets}

        wb = openpyxl.load_workbook(str(p), data_only=True)
        summary = {"total_rows": 0, "sheets": {}}

        for sheet_title in expected_sheets:
            if sheet_title not in wb.sheetnames:
                summary["sheets"][sheet_title] = 0
                continue

            ws = wb[sheet_title]
            raw_rows = []
            headers = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                if r_idx == 1:
                    headers = [str(c or '').strip() for c in row]
                    continue
                if not any(row):
                    continue
                row_dict = {}
                for col_idx, val in enumerate(row):
                    if col_idx < len(headers):
                        row_dict[headers[col_idx]] = val
                raw_rows.append(row_dict)

            rows_out = []
            for idx, r in enumerate(raw_rows):
                if sheet_title in ['Fabric Stock', 'Fabric WIP']:
                    parsed_row = {
                        'Fabric Name': str(r.get('Fabric Name') or r.get('fabric_name') or r.get('Fabric') or '').strip(),
                        'GSM': str(r.get('GSM') or r.get('gsm') or '').strip(),
                        'DIA': str(r.get('DIA') or r.get('dia') or '').strip(),
                        'Color': str(r.get('Color') or r.get('color') or '').strip(),
                        'Weight': str(r.get('Weight') or r.get('weight') or '').strip()
                    }
                elif sheet_title == 'Production WIP':
                    parsed_row = {
                        'Product Name': str(r.get('Product Name') or r.get('product_name') or '').strip(),
                        'Color': str(r.get('Color') or r.get('color') or '').strip(),
                        'Size': str(r.get('Size') or r.get('size') or '').strip(),
                        'Production Type': str(r.get('Production Type') or r.get('production_type') or 'Common').strip(),
                        'Production Group': str(r.get('Production Group') or r.get('production_group') or 'UNFOLDING').strip(),
                        'Qty': str(r.get('Qty') or r.get('qty') or '').strip()
                    }
                else: # Pending Orders, Finished Goods
                    parsed_row = {
                        'Product Name': str(r.get('Product Name') or r.get('product_name') or '').strip(),
                        'Color': str(r.get('Color') or r.get('color') or '').strip(),
                        'Size': str(r.get('Size') or r.get('size') or '').strip(),
                        'Qty': str(r.get('Qty') or r.get('qty') or '').strip()
                    }

                parsed_row['_excelRow'] = idx + 2
                parsed_row['_original'] = dict(parsed_row)
                parsed_row['_isEdited'] = False
                parsed_row['validation_status'] = 'NOT_VALIDATED'
                parsed_row['validation_message'] = 'Loaded from ERP mainout.xlsx'
                rows_out.append(parsed_row)

            parsed_sheets[sheet_title] = rows_out
            summary["sheets"][sheet_title] = len(rows_out)
            summary["total_rows"] += len(rows_out)

        return parsed_sheets, summary

    def start_sync(self, mode="review", download_dir=None, headless=False, custom_db_url=None, feed_strategy="replace"):
        """Starts batch download in background thread."""
        if download_dir:
            download_dir = str(download_dir).strip().strip('"').strip("'").strip()

        with self.lock:
            if self.job_state["status"] == "running":
                return False, "A sync process is already running."

            self.stop_event.clear()
            self.job_state = {
                "status": "running",
                "mode": mode,
                "feed_strategy": feed_strategy,
                "step_index": 0,
                "total_steps": 8,
                "step_name": "Initializing ERP modules...",
                "progress_percent": 0,
                "logs": [],
                "error": None,
                "output_file": None,
                "summary": {},
                "started_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "completed_at": None
            }

        self.worker_thread = threading.Thread(
            target=self._worker_execute,
            args=(mode, download_dir, headless, custom_db_url, feed_strategy),
            daemon=True
        )
        self.worker_thread.start()
        return True, "Sync process started."

    def _worker_execute(self, mode, download_dir, headless, custom_db_url, feed_strategy="replace"):
        try:
            self._log("=" * 60)
            self._log(f"STARTING ERP AUTO SYNC (Mode: {'Option A: Review' if mode == 'review' else 'Option B: Auto-Save'} | Feed Strategy: {'Replace (Delete Old)' if feed_strategy == 'replace' else 'Append (Keep Old)'})")
            self._log("=" * 60)

            # Ensure auto download directory is in sys.path
            source_dir_str = str(AUTO_DOWNLOAD_DIR.resolve())
            if source_dir_str not in sys.path:
                sys.path.insert(0, source_dir_str)

            # Read credentials from source .env
            env_file = AUTO_DOWNLOAD_DIR / ".env"
            env_vars = {}
            if env_file.exists():
                with open(env_file, "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line and not line.startswith("#") and "=" in line:
                            k, v = line.split("=", 1)
                            env_vars[k.strip()] = v.strip().strip('"').strip("'")

            user = env_vars.get("ERP_USERNAME")
            pwd = env_vars.get("ERP_PASSWORD")
            raw_base_dir = download_dir or env_vars.get("DOWNLOAD_DIR", r"C:\ERP_DOWNLOADS")
            base_dir = str(raw_base_dir).strip().strip('"').strip("'").strip()

            if not user or not pwd:
                raise ValueError("ERP_USERNAME or ERP_PASSWORD missing in auto download .env configuration.")

            Path(base_dir).mkdir(parents=True, exist_ok=True)
            self._log(f"Destination Base Folder: {base_dir}")
            self._log(f"Browser Automation: {'Headless' if headless else 'Visible'}")

            # Import engines dynamically
            from fabric_batch_engine import run_batch_download as run_fabric_batch
            from production_wip_engine import run_production_wip_download
            from finished_goods_engine import run_finished_goods_download
            from unfolding_stock_engine import run_unfolding_stock_download
            from cutting_received_engine import run_cutting_received_download
            from pending_order_quantity_engine import run_pending_order_quantity_download
            from fabric_orders_engine import run_fabric_orders_download
            from report_alignment_engine import run_report_alignment, load_filter_config

            # Load active fabric and WIP lists from configured settings
            all_fab = self.get_fabric_list()
            fab_items = [it for it in all_fab if it.get("enabled", True)]

            all_wip = self.get_wip_list()
            wip_items = [it for it in all_wip if it.get("enabled", True)]

            filter_cfg = self.get_filter_config()

            # -------------------------------------------------------------
            # STEP 1: Fabric Stock Batch
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 1
                self.job_state["step_name"] = "Step 1/8: Downloading Fabric Stock Reports..."
                self.job_state["progress_percent"] = 10

            self._log(f"\n>>> [1/8] DOWNLOADING FABRIC STOCK BATCH ({len(fab_items)} fabric types) <<<")
            run_fabric_batch(
                username=user,
                password=pwd,
                fabric_items=fab_items,
                download_dir=base_dir,
                headless=headless,
                status_callback=self._log,
                stop_event=self.stop_event
            )

            # -------------------------------------------------------------
            # STEP 2: Production WIP Batch
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 2
                self.job_state["step_name"] = "Step 2/8: Downloading Production WIP Reports..."
                self.job_state["progress_percent"] = 25

            self._log(f"\n>>> [2/8] DOWNLOADING PRODUCTION WIP BATCH ({len(wip_items)} groups) <<<")
            run_production_wip_download(
                username=user,
                password=pwd,
                grouping_items=wip_items,
                download_dir=base_dir,
                headless=headless,
                status_callback=self._log,
                stop_event=self.stop_event
            )

            # -------------------------------------------------------------
            # STEP 3: Finished Goods Stock
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 3
                self.job_state["step_name"] = "Step 3/8: Downloading Finished Goods Stock..."
                self.job_state["progress_percent"] = 40

            self._log("\n>>> [3/8] DOWNLOADING FINISHED GOODS STOCK <<<")
            run_finished_goods_download(
                username=user,
                password=pwd,
                download_dir=base_dir,
                headless=headless,
                status_callback=self._log,
                stop_event=self.stop_event
            )

            # -------------------------------------------------------------
            # STEP 4: Unfolding Stock
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 4
                self.job_state["step_name"] = "Step 4/8: Downloading Unfolding Stock..."
                self.job_state["progress_percent"] = 50

            self._log("\n>>> [4/8] DOWNLOADING UNFOLDING STOCK <<<")
            run_unfolding_stock_download(
                username=user,
                password=pwd,
                download_dir=base_dir,
                headless=headless,
                status_callback=self._log,
                stop_event=self.stop_event
            )

            # -------------------------------------------------------------
            # STEP 5: Cutting Received Stock
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 5
                self.job_state["step_name"] = "Step 5/8: Downloading Cutting Received Stock..."
                self.job_state["progress_percent"] = 65

            self._log("\n>>> [5/8] DOWNLOADING CUTTING RECEIVED STOCK <<<")
            run_cutting_received_download(
                username=user,
                password=pwd,
                download_dir=base_dir,
                headless=headless,
                status_callback=self._log,
                stop_event=self.stop_event
            )

            # -------------------------------------------------------------
            # STEP 6: Pending Order Quantity
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 6
                self.job_state["step_name"] = "Step 6/8: Downloading Pending Order Quantity..."
                self.job_state["progress_percent"] = 75

            self._log("\n>>> [6/8] DOWNLOADING PENDING ORDER QUANTITY <<<")
            run_pending_order_quantity_download(
                username=user,
                password=pwd,
                download_dir=base_dir,
                headless=headless,
                status_callback=self._log,
                stop_event=self.stop_event
            )

            # -------------------------------------------------------------
            # STEP 7: Fabric Orders (Pending)
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 7
                self.job_state["step_name"] = "Step 7/8: Downloading Fabric Orders (Pending)..."
                self.job_state["progress_percent"] = 85

            self._log("\n>>> [7/8] DOWNLOADING FABRIC ORDERS (PENDING) <<<")
            db_url_to_use = custom_db_url or os.getenv("GARMENT_ERP_DATABASE_URL")
            run_fabric_orders_download(
                db_url=db_url_to_use,
                download_dir=base_dir,
                status_filter="Pending",
                status_callback=self._log,
                stop_event=self.stop_event
            )

            # -------------------------------------------------------------
            # STEP 8: Auto-Compile & Align Reports into mainout.xlsx
            # -------------------------------------------------------------
            if self.stop_event.is_set():
                raise InterruptedError("Cancelled by user.")

            with self.lock:
                self.job_state["step_index"] = 8
                self.job_state["step_name"] = "Step 8/8: Compiling & Aligning Reports into mainout.xlsx..."
                self.job_state["progress_percent"] = 92

            self._log("\n>>> [8/8] COMPILING & ALIGNING REPORTS INTO MAINOUNT.XLSX <<<")
            active_filter_cfg = self.get_filter_config()
            self._log(f"Applying Sheet-wise Whitelist Filter (Enabled: {active_filter_cfg.get('enabled', True)})...")
            align_success, out_file = run_report_alignment(
                base_dir=base_dir,
                filter_config=active_filter_cfg,
                status_callback=self._log
            )

            if not align_success or not os.path.exists(out_file):
                raise RuntimeError(f"Report alignment failed or {out_file} was not generated.")

            self._log(f"[SUCCESS] mainout.xlsx generated: {out_file}")

            # Parse mainout into Bulk Feed structure
            self._log("Parsing mainout.xlsx records for Bulk Feed...")
            parsed_sheets, summary = self.parse_mainout_file(out_file)

            with self.lock:
                self.job_state["status"] = "completed"
                self.job_state["step_index"] = 8
                self.job_state["step_name"] = "Completed successfully!"
                self.job_state["progress_percent"] = 100
                self.job_state["output_file"] = out_file
                self.job_state["summary"] = summary
                self.job_state["completed_at"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            self._log(f"All 8 ERP download & alignment steps completed! Total rows: {summary.get('total_rows', 0):,}")

        except InterruptedError:
            self._log("[CANCELLED] Process cancelled by user.")
            with self.lock:
                self.job_state["status"] = "cancelled"
                self.job_state["step_name"] = "Cancelled by user"
                self.job_state["completed_at"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        except Exception as e:
            err_msg = str(e)
            self._log(f"[ERROR] Sync failed: {err_msg}")
            with self.lock:
                self.job_state["status"] = "error"
                self.job_state["error"] = err_msg
                self.job_state["step_name"] = f"Error: {err_msg}"
                self.job_state["completed_at"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# Global service instance
erp_sync_service = ErpAutoSyncService()
